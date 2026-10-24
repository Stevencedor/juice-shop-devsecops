#!/usr/bin/env python3
"""Genera docs/02_quality_gates.xlsx desde security/quality-gates.yml
(usa la plantilla del Anexo 3 si se pasa con --template).
Uso: python scripts/export_gates_xlsx.py [--template quality_gates_vacia.xlsx] [--out docs/02_quality_gates.xlsx]
"""
import argparse
from copy import copy
from pathlib import Path

import yaml
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BASE = ["ID", "Control", "Herramienta", "Métrica", "Umbral", "Acción", "Severidad", "Justificación", "Validación"]
EXTRA = ["Amenazas Act. 2", "Responsable", "Evidencia"]


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--policy", default="security/quality-gates.yml")
    ap.add_argument("--template")
    ap.add_argument("--out", default="docs/02_quality_gates.xlsx")
    a = ap.parse_args()
    gates = yaml.safe_load(Path(a.policy).read_text(encoding="utf-8"))["gates"]

    if a.template:
        wb = load_workbook(a.template)
        ws = wb.active
        hdr = ws.cell(1, 1)
        for i, name in enumerate(EXTRA, start=len(BASE) + 1):
            c = ws.cell(1, i, name)
            c.font, c.fill, c.alignment, c.border = copy(hdr.font), copy(hdr.fill), copy(hdr.alignment), copy(hdr.border)
    else:
        wb = Workbook()
        ws = wb.active
        ws.title = "Quality Gates"
        for i, name in enumerate(BASE + EXTRA, start=1):
            c = ws.cell(1, i, name)
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = PatternFill("solid", fgColor="1F3864")

    for r, g in enumerate(gates, start=2):
        row = [g["id"], g["control"], {"trivy-sca": "Trivy (fs)", "trivy-image": "Trivy (imagen)", "trivy-iac": "Trivy (config)",
                                       "zap": "OWASP ZAP", "sbom": "Syft (CycloneDX)", "gitleaks": "Gitleaks",
                                       "semgrep": "Semgrep", "exceptions": "quality_gate.py"}[g["tool"]],
               g["metric"], g["threshold"], g["action_text"], g["severity_label"],
               " ".join(g["justification"].split()), g["validation"], g["risks"], g["owner"], g["evidence"]]
        for i, v in enumerate(row, start=1):
            c = ws.cell(r, i, v)
            c.alignment = Alignment(wrap_text=True, vertical="top")
    widths = [8, 20, 18, 34, 30, 22, 11, 55, 40, 24, 22, 28]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"
    Path(a.out).parent.mkdir(parents=True, exist_ok=True)
    wb.save(a.out)
    print(f"{len(gates)} gates -> {a.out}")


if __name__ == "__main__":
    main()
